"""Real Artifact Generator — creates genuine binary documents in the workspace.

Produces valid, fully formatted files directly on disk:
- PDF documents (.pdf) via reportlab (with PDF 1.4 pure-Python fallback)
- Word documents (.docx) via python-docx (with OpenXML WordprocessingML pure-Python fallback)
- PowerPoint presentations (.pptx) via python-pptx (with OpenXML PresentationML pure-Python fallback)
- Excel spreadsheets (.xlsx) via openpyxl (with OpenXML SpreadsheetML pure-Python fallback)
- CSV data files (.csv) via standard csv module
"""

from __future__ import annotations

import csv
import io
import json
import re
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path
from typing import Any

from synapse.logging import get_logger

log = get_logger("synapse.workspace.artifacts")


# ===========================================================================
# 1. Real PDF Generator (ReportLab with pure-Python PDF 1.4 Fallback)
# ===========================================================================

def generate_pdf(title: str, text_or_markdown: str, author: str = "Synapse One") -> bytes:
    """Generate a valid, fully formatted PDF binary file."""
    try:
        return _generate_pdf_reportlab(title, text_or_markdown, author=author)
    except Exception as exc:
        log.warning("reportlab_generation_failed_falling_back", error=str(exc))
        return _generate_pdf_pure(title, text_or_markdown, author=author)


def _escape_reportlab_text(text: str) -> str:
    """Escape XML special characters for reportlab Paragraph."""
    return str(text).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _format_inline_markdown(text: str) -> str:
    """Convert basic markdown bold, italic, and inline code to reportlab XML tags."""
    s = _escape_reportlab_text(text)
    s = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", s)
    s = re.sub(r"\*(.+?)\*", r"<i>\1</i>", s)
    s = re.sub(r"(?<!\w)_(.+?)_(?!\w)", r"<i>\1</i>", s)
    s = re.sub(r"`(.+?)`", r'<font face="Courier" color="#0F172A">\1</font>', s)
    return s


def _generate_pdf_reportlab(title: str, text_or_markdown: str, author: str = "Synapse One") -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.platypus import (
        HRFlowable,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=letter,
        leftMargin=54,
        rightMargin=54,
        topMargin=54,
        bottomMargin=54,
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "CustomTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=22,
        leading=26,
        textColor=colors.HexColor("#1E293B"),
        alignment=0,
        spaceAfter=6,
    )

    author_style = ParagraphStyle(
        "CustomAuthor",
        parent=styles["Normal"],
        fontName="Helvetica-Oblique",
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#64748B"),
        spaceAfter=12,
    )

    h1_style = ParagraphStyle(
        "CustomH1",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=14,
        leading=18,
        textColor=colors.HexColor("#1E3A8A"),
        spaceBefore=12,
        spaceAfter=6,
        keepWithNext=True,
    )

    h2_style = ParagraphStyle(
        "CustomH2",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=12,
        leading=16,
        textColor=colors.HexColor("#334155"),
        spaceBefore=10,
        spaceAfter=4,
        keepWithNext=True,
    )

    body_style = ParagraphStyle(
        "CustomBody",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#334155"),
        spaceAfter=6,
    )

    bullet_style = ParagraphStyle(
        "CustomBullet",
        parent=body_style,
        leftIndent=16,
        firstLineIndent=-10,
        spaceAfter=3,
    )

    table_cell_style = ParagraphStyle(
        "TableCell",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#1E293B"),
    )

    table_header_style = ParagraphStyle(
        "TableHeader",
        parent=table_cell_style,
        fontName="Helvetica-Bold",
        textColor=colors.white,
    )

    story = []

    if title:
        story.append(Paragraph(_escape_reportlab_text(title), title_style))
        story.append(Paragraph(f"Author: {_escape_reportlab_text(author)}", author_style))
        story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#CBD5E1"), spaceBefore=0, spaceAfter=12))

    lines = text_or_markdown.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue

        # Markdown table detection
        if "|" in line:
            table_lines = [line]
            i += 1
            while i < len(lines) and "|" in lines[i]:
                table_lines.append(lines[i].strip())
                i += 1

            raw_table_rows = []
            for tl in table_lines:
                if re.match(r"^\|?\s*[-:]+\s*\|", tl):
                    continue
                cells = [c.strip() for c in tl.strip("|").split("|")]
                if cells:
                    raw_table_rows.append(cells)

            if raw_table_rows:
                table_data = []
                for r_idx, row in enumerate(raw_table_rows):
                    row_data = []
                    style_to_use = table_header_style if r_idx == 0 else table_cell_style
                    for c in row:
                        row_data.append(Paragraph(_format_inline_markdown(c), style_to_use))
                    table_data.append(row_data)

                t = Table(table_data, hAlign="LEFT")
                t.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E3A8A")),
                    ("ALIGN", (0, 0), (-1, -1), "LEFT"),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                    ("LEFTPADDING", (0, 0), (-1, -1), 6),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
                ]))
                story.append(t)
                story.append(Spacer(1, 8))
            continue

        if line.startswith("### "):
            story.append(Paragraph(_format_inline_markdown(line[4:]), h2_style))
        elif line.startswith("## "):
            story.append(Paragraph(_format_inline_markdown(line[3:]), h1_style))
        elif line.startswith("# "):
            story.append(Paragraph(_format_inline_markdown(line[2:]), h1_style))
        elif line.startswith(("- ", "* ", "• ")):
            bullet_text = "&bull; " + _format_inline_markdown(line[2:].strip())
            story.append(Paragraph(bullet_text, bullet_style))
        elif re.match(r"^\d+\.\s+", line):
            m = re.match(r"^(\d+\.)\s+(.*)", line)
            if m:
                num_text = f"<b>{m.group(1)}</b> " + _format_inline_markdown(m.group(2))
                story.append(Paragraph(num_text, bullet_style))
            else:
                story.append(Paragraph(_format_inline_markdown(line), body_style))
        elif line.startswith("---") or line.startswith("***"):
            story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#E2E8F0"), spaceBefore=6, spaceAfter=8))
        else:
            story.append(Paragraph(_format_inline_markdown(line), body_style))
        i += 1

    if not story:
        story.append(Paragraph(_escape_reportlab_text(title or "Document"), title_style))
        story.append(Paragraph("Generated by Synapse One", body_style))

    doc.build(story)
    return buf.getvalue()


def _generate_pdf_pure(title: str, text_or_markdown: str, author: str = "Synapse One") -> bytes:
    """Generate a valid PDF 1.4 binary file without third-party dependencies."""
    lines = _format_text_lines(text_or_markdown)

    page_width, page_height = 612.0, 792.0  # US Letter
    margin_x, margin_y = 54.0, 54.0
    usable_height = page_height - (2 * margin_y)

    pages_lines: list[list[tuple[str, str]]] = []
    current_page: list[tuple[str, str]] = [("title", title)]
    current_y = usable_height - 50.0

    for kind, text in lines:
        line_height = 24.0 if kind == "heading" else 14.0
        if current_y - line_height < 0:
            pages_lines.append(current_page)
            current_page = []
            current_y = usable_height
        current_page.append((kind, text))
        current_y -= line_height

    if current_page:
        pages_lines.append(current_page)

    if not pages_lines:
        pages_lines = [[("title", title), ("body", "Generated by Synapse One")]]

    objects: list[bytes] = []

    def add_object(content: bytes) -> int:
        objects.append(content)
        return len(objects)

    add_object(b"<< /Type /Catalog /Pages 2 0 R >>")
    pages_obj_idx = 1
    add_object(b"")
    font_reg_id = add_object(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    font_bold_id = add_object(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold >>")

    page_obj_ids = []

    for p_idx, page_content in enumerate(pages_lines, 1):
        stream_lines = []
        stream_lines.append(b"BT")

        y = page_height - margin_y
        for kind, text in page_content:
            safe_text = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            if kind == "title":
                y -= 28.0
                stream_lines.append(f"/F2 18 Tf {margin_x:.2f} {y:.2f} Td ({safe_text}) Tj".encode("latin-1", "replace"))
            elif kind == "heading":
                y -= 20.0
                stream_lines.append(f"/F2 13 Tf {margin_x:.2f} {y:.2f} Td ({safe_text}) Tj".encode("latin-1", "replace"))
            else:
                y -= 14.0
                stream_lines.append(f"/F1 10 Tf {margin_x:.2f} {y:.2f} Td ({safe_text}) Tj".encode("latin-1", "replace"))

        stream_lines.append(b"ET")
        stream_data = b"\n".join(stream_lines)

        contents_id = add_object(
            f"<< /Length {len(stream_data)} >>\nstream\n".encode("ascii")
            + stream_data
            + b"\nendstream"
        )

        page_id = add_object(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 {page_width} {page_height}] "
            f"/Resources << /Font << /F1 {font_reg_id} 0 R /F2 {font_bold_id} 0 R >> >> "
            f"/Contents {contents_id} 0 R >>".encode("ascii")
        )
        page_obj_ids.append(page_id)

    kids_str = " ".join(f"{pid} 0 R" for pid in page_obj_ids)
    objects[pages_obj_idx] = f"<< /Type /Pages /Kids [{kids_str}] /Count {len(page_obj_ids)} >>".encode("ascii")

    buf = io.BytesIO()
    buf.write(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0]

    for i, obj in enumerate(objects, 1):
        offsets.append(buf.tell())
        buf.write(f"{i} 0 obj\n".encode("ascii"))
        buf.write(obj)
        buf.write(b"\nendobj\n")

    xref_offset = buf.tell()
    buf.write(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    buf.write(b"0000000000 65535 f \n")
    for off in offsets[1:]:
        buf.write(f"{off:010d} 00000 n \n".encode("ascii"))

    buf.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n".encode("ascii"))
    return buf.getvalue()


def _format_text_lines(raw: str) -> list[tuple[str, str]]:
    """Convert raw markdown / text into wrapped line items: [(kind, text)]."""
    result: list[tuple[str, str]] = []
    for raw_line in raw.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("#"):
            title_text = line.lstrip("#").strip()
            result.append(("heading", title_text))
        elif line.startswith(("-", "*", "•")):
            item_text = "• " + line.lstrip("-*•").strip()
            result.extend([("body", sub) for sub in _wrap_text(item_text, 75)])
        else:
            result.extend([("body", sub) for sub in _wrap_text(line, 80)])
    return result


def _wrap_text(text: str, width: int) -> list[str]:
    words = text.split()
    if not words:
        return []
    lines = []
    curr: list[str] = []
    curr_len = 0
    for w in words:
        if curr_len + len(w) + (1 if curr else 0) > width:
            lines.append(" ".join(curr))
            curr = [w]
            curr_len = len(w)
        else:
            curr.append(w)
            curr_len += len(w) + (1 if len(curr) > 1 else 0)
    if curr:
        lines.append(" ".join(curr))
    return lines


# ===========================================================================
# 2. Real Word Generator (.docx)
# ===========================================================================

def generate_docx(title: str, text_or_markdown: str, author: str = "Synapse One") -> bytes:
    """Generate a valid, fully formatted Microsoft Word (.docx) binary."""
    try:
        return _generate_docx_library(title, text_or_markdown, author=author)
    except Exception as exc:
        log.warning("docx_library_failed_falling_back", error=str(exc))
        return _generate_docx_pure(title, text_or_markdown, author=author)


def _generate_docx_library(title: str, text_or_markdown: str, author: str = "Synapse One") -> bytes:
    import docx
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.shared import RGBColor

    doc = docx.Document()

    if title:
        doc.add_heading(title, level=0)
        p_author = doc.add_paragraph(f"Author: {author}")
        if p_author.runs:
            p_author.runs[0].font.italic = True
            p_author.runs[0].font.color.rgb = RGBColor(100, 116, 139)

    lines = text_or_markdown.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue

        if "|" in line:
            table_lines = [line]
            i += 1
            while i < len(lines) and "|" in lines[i]:
                table_lines.append(lines[i].strip())
                i += 1

            raw_table_rows = []
            for tl in table_lines:
                if re.match(r"^\|?\s*[-:]+\s*\|", tl):
                    continue
                cells = [c.strip() for c in tl.strip("|").split("|")]
                if cells:
                    raw_table_rows.append(cells)

            if raw_table_rows:
                cols_count = max(len(r) for r in raw_table_rows)
                t = doc.add_table(rows=len(raw_table_rows), cols=cols_count)
                t.alignment = WD_TABLE_ALIGNMENT.LEFT
                t.style = 'Table Grid'
                for r_idx, row in enumerate(raw_table_rows):
                    for c_idx, cell_text in enumerate(row):
                        if c_idx < cols_count:
                            cell = t.cell(r_idx, c_idx)
                            cell.text = cell_text
                            if r_idx == 0:
                                for p in cell.paragraphs:
                                    for r in p.runs:
                                        r.font.bold = True
                doc.add_paragraph()
            continue

        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
        elif line.startswith(("- ", "* ", "• ")):
            p = doc.add_paragraph(style='List Bullet')
            _add_formatted_runs_docx(p, line[2:].strip())
        elif re.match(r"^\d+\.\s+", line):
            m = re.match(r"^\d+\.\s+(.*)", line)
            p = doc.add_paragraph(style='List Number')
            _add_formatted_runs_docx(p, m.group(1).strip() if m else line)
        else:
            p = doc.add_paragraph()
            _add_formatted_runs_docx(p, line)
        i += 1

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _add_formatted_runs_docx(paragraph, text: str) -> None:
    tokens = re.split(r"(\*\*.*?\*\*|\*.*?\*|`.*?`)", text)
    for token in tokens:
        if not token:
            continue
        if token.startswith("**") and token.endswith("**") and len(token) >= 4:
            r = paragraph.add_run(token[2:-2])
            r.bold = True
        elif token.startswith("*") and token.endswith("*") and len(token) >= 2:
            r = paragraph.add_run(token[1:-1])
            r.italic = True
        elif token.startswith("`") and token.endswith("`") and len(token) >= 2:
            r = paragraph.add_run(token[1:-1])
            r.font.name = "Consolas"
        else:
            paragraph.add_run(token)


def _generate_docx_pure(title: str, text_or_markdown: str, author: str = "Synapse One") -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        ct = ET.Element("Types", xmlns="http://schemas.openxmlformats.org/package/2006/content-types")
        ET.SubElement(ct, "Default", Extension="rels", ContentType="application/vnd.openxmlformats-package.relationships+xml")
        ET.SubElement(ct, "Default", Extension="xml", ContentType="application/xml")
        ET.SubElement(ct, "Override", PartName="/word/document.xml", ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml")
        z.writestr("[Content_Types].xml", ET.tostring(ct, encoding="utf-8", xml_declaration=True))

        rels = ET.Element("Relationships", xmlns="http://schemas.openxmlformats.org/package/2006/relationships")
        ET.SubElement(rels, "Relationship", Id="rId1", Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument", Target="word/document.xml")
        z.writestr("_rels/.rels", ET.tostring(rels, encoding="utf-8", xml_declaration=True))

        w_doc = ET.Element("w:document", {
            "xmlns:w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
            "xmlns:r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
        })
        body = ET.SubElement(w_doc, "w:body")

        if title:
            p_title = ET.SubElement(body, "w:p")
            pPr = ET.SubElement(p_title, "w:pPr")
            rPr = ET.SubElement(pPr, "w:rPr")
            ET.SubElement(rPr, "w:b")
            ET.SubElement(rPr, "w:sz", {"w:val": "36"})
            r_title = ET.SubElement(p_title, "w:r")
            t_title = ET.SubElement(r_title, "w:t")
            t_title.text = title

        for line in text_or_markdown.splitlines():
            line_str = line.strip()
            if not line_str:
                continue
            p = ET.SubElement(body, "w:p")
            if line_str.startswith("#"):
                clean = line_str.lstrip("#").strip()
                pPr = ET.SubElement(p, "w:pPr")
                rPr = ET.SubElement(pPr, "w:rPr")
                ET.SubElement(rPr, "w:b")
                ET.SubElement(rPr, "w:sz", {"w:val": "28"})
                r = ET.SubElement(p, "w:r")
                t = ET.SubElement(r, "w:t")
                t.text = clean
            else:
                r = ET.SubElement(p, "w:r")
                t = ET.SubElement(r, "w:t")
                t.text = line_str

        z.writestr("word/document.xml", ET.tostring(w_doc, encoding="utf-8", xml_declaration=True))

    return buf.getvalue()


# ===========================================================================
# 3. Real PPTX Generator (PresentationML with python-pptx & fallback)
# ===========================================================================

def generate_pptx(title: str, slides_content: list[tuple[str, list[str]]] | list[dict] | str) -> bytes:
    """Generate a valid, native Microsoft PowerPoint (.pptx) presentation binary."""
    try:
        return _generate_pptx_library(title, slides_content)
    except Exception as exc:
        log.warning("pptx_library_failed_falling_back", error=str(exc))
        return _generate_pptx_pure(title, slides_content)


def _normalize_slides(main_title: str, slides_content: list[tuple[str, list[str]]] | list[dict] | str) -> list[tuple[str, list[str]]]:
    """Normalize slide structures from dicts, tuples, or raw markdown text."""
    if isinstance(slides_content, list):
        slides = []
        for item in slides_content:
            if isinstance(item, tuple) and len(item) == 2:
                title, bullets = item
                slides.append((str(title), [str(b) for b in bullets]))
            elif isinstance(item, dict):
                title = item.get("title") or item.get("name") or "Slide"
                bullets = item.get("bullets") or item.get("content") or item.get("points") or []
                if isinstance(bullets, str):
                    bullets = [b.strip("-*• ") for b in bullets.splitlines() if b.strip()]
                slides.append((str(title), [str(b) for b in bullets]))
        return slides
    if isinstance(slides_content, str):
        try:
            parsed = json.loads(slides_content)
            if isinstance(parsed, list):
                return _normalize_slides(main_title, parsed)
            if isinstance(parsed, dict) and "slides" in parsed and isinstance(parsed["slides"], list):
                return _normalize_slides(main_title, parsed["slides"])
        except Exception:
            pass
        return _parse_slides_from_text(main_title, slides_content)
    return []


def _generate_pptx_library(title: str, slides_content: list[tuple[str, list[str]]] | list[dict] | str) -> bytes:
    import pptx
    from pptx.util import Inches, Pt

    prs = pptx.Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slides = _normalize_slides(title, slides_content)

    # Title slide
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    slide.shapes.title.text = title or "Presentation"
    if len(slide.placeholders) > 1:
        slide.placeholders[1].text = "Generated by Synapse One"

    # Content slides
    content_layout = prs.slide_layouts[1]
    for slide_title, bullets in slides:
        s = prs.slides.add_slide(content_layout)
        s.shapes.title.text = slide_title
        if len(s.placeholders) > 1:
            tf = s.placeholders[1].text_frame
            tf.clear()
            for idx, b in enumerate(bullets):
                p = tf.add_paragraph() if idx > 0 else tf.paragraphs[0]
                p.text = str(b)
                p.level = 0
                p.font.size = Pt(18)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _generate_pptx_pure(title: str, slides_content: list[tuple[str, list[str]]] | list[dict] | str) -> bytes:
    slides = _normalize_slides(title, slides_content)
    if not slides:
        slides = [(title or "Presentation", [])]

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        ct = ET.Element("Types", xmlns="http://schemas.openxmlformats.org/package/2006/content-types")
        ET.SubElement(ct, "Default", Extension="rels", ContentType="application/vnd.openxmlformats-package.relationships+xml")
        ET.SubElement(ct, "Default", Extension="xml", ContentType="application/xml")
        ET.SubElement(ct, "Override", PartName="/ppt/presentation.xml", ContentType="application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml")
        ET.SubElement(ct, "Override", PartName="/ppt/slideMasters/slideMaster1.xml", ContentType="application/vnd.openxmlformats-officedocument.presentationml.slideMaster+xml")
        ET.SubElement(ct, "Override", PartName="/ppt/slideLayouts/slideLayout1.xml", ContentType="application/vnd.openxmlformats-officedocument.presentationml.slideLayout+xml")

        for i in range(1, len(slides) + 1):
            ET.SubElement(ct, "Override", PartName=f"/ppt/slides/slide{i}.xml", ContentType="application/vnd.openxmlformats-officedocument.presentationml.slide+xml")
        z.writestr("[Content_Types].xml", ET.tostring(ct, encoding="utf-8", xml_declaration=True))

        rels = ET.Element("Relationships", xmlns="http://schemas.openxmlformats.org/package/2006/relationships")
        ET.SubElement(rels, "Relationship", Id="rId1", Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument", Target="ppt/presentation.xml")
        z.writestr("_rels/.rels", ET.tostring(rels, encoding="utf-8", xml_declaration=True))

        p_rels = ET.Element("Relationships", xmlns="http://schemas.openxmlformats.org/package/2006/relationships")
        ET.SubElement(p_rels, "Relationship", Id="rId1", Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideMaster", Target="slideMasters/slideMaster1.xml")
        for i in range(1, len(slides) + 1):
            ET.SubElement(p_rels, "Relationship", Id=f"rId{i+1}", Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide", Target=f"slides/slide{i}.xml")
        z.writestr("ppt/_rels/presentation.xml.rels", ET.tostring(p_rels, encoding="utf-8", xml_declaration=True))

        p = ET.Element("p:presentation", {
            "xmlns:a": "http://schemas.openxmlformats.org/drawingml/2006/main",
            "xmlns:r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
            "xmlns:p": "http://schemas.openxmlformats.org/presentationml/2006/main"
        })
        sldMasterIdLst = ET.SubElement(p, "p:sldMasterIdLst")
        ET.SubElement(sldMasterIdLst, "p:sldMasterId", {"id": "2147483648", "r:id": "rId1"})
        sldIdLst = ET.SubElement(p, "p:sldIdLst")
        for i in range(1, len(slides) + 1):
            ET.SubElement(sldIdLst, "p:sldId", {"id": str(255 + i), "r:id": f"rId{i+1}"})
        ET.SubElement(p, "p:sldSz", {"cx": "9144000", "cy": "6858000", "type": "screen4x3"})
        z.writestr("ppt/presentation.xml", ET.tostring(p, encoding="utf-8", xml_declaration=True))

        _write_pptx_master_and_layout(z)

        for i, (slide_title, bullets) in enumerate(slides, 1):
            _write_pptx_slide(z, i, slide_title, bullets)

    return buf.getvalue()


def _write_pptx_master_and_layout(z: zipfile.ZipFile) -> None:
    sm = ET.Element("p:sldMaster", {
        "xmlns:a": "http://schemas.openxmlformats.org/drawingml/2006/main",
        "xmlns:r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
        "xmlns:p": "http://schemas.openxmlformats.org/presentationml/2006/main"
    })
    cSld = ET.SubElement(sm, "p:cSld")
    spTree = ET.SubElement(cSld, "p:spTree")
    nvGrpSpPr = ET.SubElement(spTree, "p:nvGrpSpPr")
    ET.SubElement(nvGrpSpPr, "p:cNvPr", {"id": "1", "name": ""})
    ET.SubElement(nvGrpSpPr, "p:cNvGrpSpPr")
    ET.SubElement(nvGrpSpPr, "p:nvPr")
    ET.SubElement(spTree, "p:grpSpPr")
    sldLayoutIdLst = ET.SubElement(sm, "p:sldLayoutIdLst")
    ET.SubElement(sldLayoutIdLst, "p:sldLayoutId", {"id": "2147483649", "r:id": "rId1"})
    z.writestr("ppt/slideMasters/slideMaster1.xml", ET.tostring(sm, encoding="utf-8", xml_declaration=True))

    sm_rels = ET.Element("Relationships", xmlns="http://schemas.openxmlformats.org/package/2006/relationships")
    ET.SubElement(sm_rels, "Relationship", Id="rId1", Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideLayout", Target="../slideLayouts/slideLayout1.xml")
    z.writestr("ppt/slideMasters/_rels/slideMaster1.xml.rels", ET.tostring(sm_rels, encoding="utf-8", xml_declaration=True))

    sl = ET.Element("p:sldLayout", {
        "xmlns:a": "http://schemas.openxmlformats.org/drawingml/2006/main",
        "xmlns:r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
        "xmlns:p": "http://schemas.openxmlformats.org/presentationml/2006/main",
        "type": "titleAndBody"
    })
    cSld = ET.SubElement(sl, "p:cSld", {"name": "Title and Content"})
    spTree = ET.SubElement(cSld, "p:spTree")
    nvGrpSpPr = ET.SubElement(spTree, "p:nvGrpSpPr")
    ET.SubElement(nvGrpSpPr, "p:cNvPr", {"id": "1", "name": ""})
    ET.SubElement(nvGrpSpPr, "p:cNvGrpSpPr")
    ET.SubElement(nvGrpSpPr, "p:nvPr")
    ET.SubElement(spTree, "p:grpSpPr")
    z.writestr("ppt/slideLayouts/slideLayout1.xml", ET.tostring(sl, encoding="utf-8", xml_declaration=True))

    sl_rels = ET.Element("Relationships", xmlns="http://schemas.openxmlformats.org/package/2006/relationships")
    ET.SubElement(sl_rels, "Relationship", Id="rId1", Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideMaster", Target="../slideMasters/slideMaster1.xml")
    z.writestr("ppt/slideLayouts/_rels/slideLayout1.xml.rels", ET.tostring(sl_rels, encoding="utf-8", xml_declaration=True))


def _write_pptx_slide(z: zipfile.ZipFile, index: int, title: str, bullets: list[str]) -> None:
    s = ET.Element("p:sld", {
        "xmlns:a": "http://schemas.openxmlformats.org/drawingml/2006/main",
        "xmlns:r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
        "xmlns:p": "http://schemas.openxmlformats.org/presentationml/2006/main"
    })
    cSld = ET.SubElement(s, "p:cSld")
    spTree = ET.SubElement(cSld, "p:spTree")

    nvGrpSpPr = ET.SubElement(spTree, "p:nvGrpSpPr")
    ET.SubElement(nvGrpSpPr, "p:cNvPr", {"id": "1", "name": ""})
    ET.SubElement(nvGrpSpPr, "p:cNvGrpSpPr")
    ET.SubElement(nvGrpSpPr, "p:nvPr")
    ET.SubElement(spTree, "p:grpSpPr")

    sp_title = ET.SubElement(spTree, "p:sp")
    nvSpPr = ET.SubElement(sp_title, "p:nvSpPr")
    ET.SubElement(nvSpPr, "p:cNvPr", {"id": "2", "name": "Title 1"})
    ET.SubElement(nvSpPr, "p:cNvSpPr")
    ET.SubElement(nvSpPr, "p:nvPr")
    spPr = ET.SubElement(sp_title, "p:spPr")
    xfrm = ET.SubElement(spPr, "a:xfrm")
    ET.SubElement(xfrm, "a:off", {"x": "685800", "y": "457200"})
    ET.SubElement(xfrm, "a:ext", {"cx": "7772400", "cy": "1143000"})
    txBody = ET.SubElement(sp_title, "p:txBody")
    ET.SubElement(txBody, "a:bodyPr")
    p_elem = ET.SubElement(txBody, "a:p")
    r_elem = ET.SubElement(p_elem, "a:r")
    ET.SubElement(r_elem, "a:rPr", {"lang": "en-US", "sz": "3200", "b": "1"})
    t_elem = ET.SubElement(r_elem, "a:t")
    t_elem.text = title

    sp_body = ET.SubElement(spTree, "p:sp")
    nvSpPr_b = ET.SubElement(sp_body, "p:nvSpPr")
    ET.SubElement(nvSpPr_b, "p:cNvPr", {"id": "3", "name": "Content Placeholder 2"})
    ET.SubElement(nvSpPr_b, "p:cNvSpPr")
    ET.SubElement(nvSpPr_b, "p:nvPr")
    spPr_b = ET.SubElement(sp_body, "p:spPr")
    xfrm_b = ET.SubElement(spPr_b, "a:xfrm")
    ET.SubElement(xfrm_b, "a:off", {"x": "685800", "y": "1828800"})
    ET.SubElement(xfrm_b, "a:ext", {"cx": "7772400", "cy": "4572000"})
    txBody_b = ET.SubElement(sp_body, "p:txBody")
    ET.SubElement(txBody_b, "a:bodyPr")

    for b in bullets:
        p_b = ET.SubElement(txBody_b, "a:p")
        ET.SubElement(p_b, "a:pPr", {"lvl": "0"})
        r_b = ET.SubElement(p_b, "a:r")
        ET.SubElement(r_b, "a:rPr", {"lang": "en-US", "sz": "1800"})
        t_b = ET.SubElement(r_b, "a:t")
        t_b.text = b

    z.writestr(f"ppt/slides/slide{index}.xml", ET.tostring(s, encoding="utf-8", xml_declaration=True))

    s_rels = ET.Element("Relationships", xmlns="http://schemas.openxmlformats.org/package/2006/relationships")
    ET.SubElement(s_rels, "Relationship", Id="rId1", Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideLayout", Target="../slideLayouts/slideLayout1.xml")
    z.writestr(f"ppt/slides/_rels/slide{index}.xml.rels", ET.tostring(s_rels, encoding="utf-8", xml_declaration=True))


def _parse_slides_from_text(main_title: str, text: str) -> list[tuple[str, list[str]]]:
    """Parse slide titles and bullet points from markdown text."""
    slides: list[tuple[str, list[str]]] = []
    current_title = main_title
    current_bullets: list[str] = []

    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        if line.startswith(("#", "Slide", "slide", "SLIDE", "---")):
            if current_bullets:
                slides.append((current_title, current_bullets))
                current_bullets = []
            clean_title = re.sub(r"^(#+|Slide\s*\d+:?|slide\s*\d+:?|---)\s*", "", line).strip()
            if clean_title:
                current_title = clean_title
        elif line.startswith(("-", "*", "•")):
            current_bullets.append(line.lstrip("-*• ").strip())
        else:
            current_bullets.append(line)

    if current_bullets or current_title:
        slides.append((current_title, current_bullets or ["Overview"]))
    return slides


# ===========================================================================
# 4. Real XLSX & CSV Generator (openpyxl & SpreadsheetML)
# ===========================================================================

def generate_csv(headers: list[str], rows: list[list[Any]]) -> str:
    """Generate standard formatted CSV content."""
    output = io.StringIO()
    writer = csv.writer(output, lineterminator="\n")
    if headers:
        writer.writerow(headers)
    for r in rows:
        writer.writerow(r)
    return output.getvalue()


def generate_xlsx(sheet_name: str, headers: list[str], rows: list[list[Any]]) -> bytes:
    """Generate a valid, native Microsoft Excel (.xlsx) spreadsheet binary."""
    try:
        return _generate_xlsx_library(sheet_name, headers, rows)
    except Exception as exc:
        log.warning("xlsx_library_failed_falling_back", error=str(exc))
        return _generate_xlsx_pure(sheet_name, headers, rows)


def _generate_xlsx_library(sheet_name: str, headers: list[str], rows: list[list[Any]]) -> bytes:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = (sheet_name or "Sheet1")[:31]

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1'),
    )

    current_row = 1
    if headers:
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=current_row, column=col_idx, value=str(h))
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
        current_row += 1

    for r_idx, row_data in enumerate(rows):
        is_alt = (r_idx % 2 == 1)
        for col_idx, val in enumerate(row_data, 1):
            val_str = str(val).strip()
            typed_val: Any = val_str
            if re.match(r"^-?\d+$", val_str):
                try:
                    typed_val = int(val_str)
                except ValueError:
                    pass
            elif re.match(r"^-?\d+\.\d+$", val_str):
                try:
                    typed_val = float(val_str)
                except ValueError:
                    pass

            cell = ws.cell(row=current_row, column=col_idx, value=typed_val)
            if is_alt:
                cell.fill = alt_fill
            cell.border = thin_border
            if isinstance(typed_val, (int, float)):
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
        current_row += 1

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _generate_xlsx_pure(sheet_name: str, headers: list[str], rows: list[list[Any]]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        ct = ET.Element("Types", xmlns="http://schemas.openxmlformats.org/package/2006/content-types")
        ET.SubElement(ct, "Default", Extension="rels", ContentType="application/vnd.openxmlformats-package.relationships+xml")
        ET.SubElement(ct, "Default", Extension="xml", ContentType="application/xml")
        ET.SubElement(ct, "Override", PartName="/xl/workbook.xml", ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml")
        ET.SubElement(ct, "Override", PartName="/xl/worksheets/sheet1.xml", ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml")
        ET.SubElement(ct, "Override", PartName="/xl/styles.xml", ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml")
        z.writestr("[Content_Types].xml", ET.tostring(ct, encoding="utf-8", xml_declaration=True))

        rels = ET.Element("Relationships", xmlns="http://schemas.openxmlformats.org/package/2006/relationships")
        ET.SubElement(rels, "Relationship", Id="rId1", Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument", Target="xl/workbook.xml")
        z.writestr("_rels/.rels", ET.tostring(rels, encoding="utf-8", xml_declaration=True))

        wb_rels = ET.Element("Relationships", xmlns="http://schemas.openxmlformats.org/package/2006/relationships")
        ET.SubElement(wb_rels, "Relationship", Id="rId1", Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet", Target="worksheets/sheet1.xml")
        ET.SubElement(wb_rels, "Relationship", Id="rId2", Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles", Target="styles.xml")
        z.writestr("xl/_rels/workbook.xml.rels", ET.tostring(wb_rels, encoding="utf-8", xml_declaration=True))

        wb = ET.Element("workbook", {"xmlns": "http://schemas.openxmlformats.org/spreadsheetml/2006/main", "xmlns:r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"})
        sheets = ET.SubElement(wb, "sheets")
        ET.SubElement(sheets, "sheet", {"name": sheet_name or "Sheet1", "sheetId": "1", "r:id": "rId1"})
        z.writestr("xl/workbook.xml", ET.tostring(wb, encoding="utf-8", xml_declaration=True))

        styles = ET.Element("styleSheet", {"xmlns": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"})
        fonts = ET.SubElement(styles, "fonts", {"count": "2"})
        f0 = ET.SubElement(fonts, "font")
        ET.SubElement(f0, "sz", {"val": "11"})
        ET.SubElement(f0, "name", {"val": "Calibri"})
        f1 = ET.SubElement(fonts, "font")
        ET.SubElement(f1, "b")
        ET.SubElement(f1, "sz", {"val": "11"})
        ET.SubElement(f1, "name", {"val": "Calibri"})

        fills = ET.SubElement(styles, "fills", {"count": "2"})
        ET.SubElement(ET.SubElement(fills, "fill"), "patternFill", {"patternType": "none"})
        ET.SubElement(ET.SubElement(fills, "fill"), "patternFill", {"patternType": "gray125"})

        borders = ET.SubElement(styles, "borders", {"count": "1"})
        ET.SubElement(borders, "border")

        cellXfs = ET.SubElement(styles, "cellXfs", {"count": "2"})
        ET.SubElement(cellXfs, "xf", {"numFmtId": "0", "fontId": "0", "fillId": "0", "borderId": "0"})
        ET.SubElement(cellXfs, "xf", {"numFmtId": "0", "fontId": "1", "fillId": "0", "borderId": "0", "applyFont": "1"})
        z.writestr("xl/styles.xml", ET.tostring(styles, encoding="utf-8", xml_declaration=True))

        ws = ET.Element("worksheet", {"xmlns": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"})
        sheetData = ET.SubElement(ws, "sheetData")

        row_idx = 1
        if headers:
            r_elem = ET.SubElement(sheetData, "row", {"r": str(row_idx)})
            for col_idx, h in enumerate(headers, 1):
                col_letter = _col_to_letter(col_idx)
                c_elem = ET.SubElement(r_elem, "c", {"r": f"{col_letter}{row_idx}", "t": "inlineStr", "s": "1"})
                is_elem = ET.SubElement(c_elem, "is")
                t_elem = ET.SubElement(is_elem, "t")
                t_elem.text = str(h)
            row_idx += 1

        for r_data in rows:
            r_elem = ET.SubElement(sheetData, "row", {"r": str(row_idx)})
            for col_idx, val in enumerate(r_data, 1):
                col_letter = _col_to_letter(col_idx)
                val_str = str(val).strip()
                if re.match(r"^-?\d+(\.\d+)?$", val_str):
                    c_elem = ET.SubElement(r_elem, "c", {"r": f"{col_letter}{row_idx}"})
                    v_elem = ET.SubElement(c_elem, "v")
                    v_elem.text = val_str
                else:
                    c_elem = ET.SubElement(r_elem, "c", {"r": f"{col_letter}{row_idx}", "t": "inlineStr"})
                    is_elem = ET.SubElement(c_elem, "is")
                    t_elem = ET.SubElement(is_elem, "t")
                    t_elem.text = val_str
            row_idx += 1

        z.writestr("xl/worksheets/sheet1.xml", ET.tostring(ws, encoding="utf-8", xml_declaration=True))

    return buf.getvalue()


def _col_to_letter(col_idx: int) -> str:
    result = ""
    while col_idx > 0:
        col_idx, remainder = divmod(col_idx - 1, 26)
        result = chr(65 + remainder) + result
    return result


def _parse_table_data(text: str) -> tuple[list[str], list[list[str]]]:
    """Extract tabular headers and rows from CSV, markdown tables, or plain lines."""
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    if not lines:
        return ["Item", "Details"], [["1", "Sample Record"]]

    if any("|" in l for l in lines):
        table_rows = []
        for l in lines:
            if re.match(r"^\|?\s*[-:]+\s*\|", l):
                continue
            cells = [c.strip() for c in l.strip("|").split("|")]
            if cells:
                table_rows.append(cells)
        if table_rows:
            return table_rows[0], table_rows[1:]

    try:
        reader = list(csv.reader(lines))
        if reader and len(reader[0]) > 1:
            return reader[0], reader[1:]
    except Exception:
        pass

    return ["Item", "Value"], [[f"Entry {i}", line] for i, line in enumerate(lines, 1)]


# ===========================================================================
# 5. Universal Artifact Dispatcher
# ===========================================================================

def build_artifact_content(rel_path: str, raw_content: str | bytes | dict | list) -> tuple[bytes | str, bool]:
    """Convert raw text / descriptions into a valid binary or formatted artifact if needed.

    Returns (data, is_binary).
    """
    ext = Path(rel_path).suffix.lower()

    if isinstance(raw_content, bytes):
        return raw_content, True

    if isinstance(raw_content, (dict, list)):
        text_content = json.dumps(raw_content)
    else:
        text_content = str(raw_content or "")

    json_obj = None
    try:
        if text_content.strip().startswith("{") and text_content.strip().endswith("}"):
            json_obj = json.loads(text_content.strip())
    except Exception:
        pass

    def _format_structured_body(data_obj: dict, default_body: str, default_title: str) -> str:
        if "questions" in data_obj and isinstance(data_obj["questions"], list):
            q_lines = [f"# {data_obj.get('title') or default_title}", ""]
            if "instructions" in data_obj:
                q_lines.append(f"**Instructions:** {data_obj['instructions']}\n")
            for idx, q in enumerate(data_obj["questions"], 1):
                if isinstance(q, dict):
                    q_text = q.get("question") or q.get("text") or q.get("prompt") or ""
                    marks = f" [{q.get('marks')} Marks]" if "marks" in q else ""
                    q_lines.append(f"**Question {idx}:** {q_text}{marks}")
                    if "options" in q and isinstance(q["options"], list):
                        for opt in q["options"]:
                            q_lines.append(f"- {opt}")
                    q_lines.append("")
                else:
                    q_lines.append(f"**Question {idx}:** {q}\n")
            return "\n".join(q_lines)
        if "sections" in data_obj and isinstance(data_obj["sections"], list):
            s_lines = [f"# {data_obj.get('title') or default_title}", ""]
            for s in data_obj["sections"]:
                if isinstance(s, dict):
                    s_title = s.get("title") or s.get("heading") or "Section"
                    s_content = s.get("content") or s.get("body") or s.get("text") or ""
                    s_lines.append(f"## {s_title}\n{s_content}\n")
                else:
                    s_lines.append(f"{s}\n")
            return "\n".join(s_lines)
        return data_obj.get("content") or data_obj.get("text") or data_obj.get("body") or default_body

    if ext == ".pdf":
        title = Path(rel_path).stem.replace("_", " ").title()
        author = "Synapse One"
        body_text = text_content
        if json_obj and isinstance(json_obj, dict):
            title = json_obj.get("title") or title
            author = json_obj.get("author") or author
            body_text = _format_structured_body(json_obj, text_content, title)
        return generate_pdf(title, body_text, author=author), True

    if ext in (".docx", ".doc"):
        title = Path(rel_path).stem.replace("_", " ").title()
        author = "Synapse One"
        body_text = text_content
        if json_obj and isinstance(json_obj, dict):
            title = json_obj.get("title") or title
            author = json_obj.get("author") or author
            body_text = _format_structured_body(json_obj, text_content, title)
        return generate_docx(title, body_text, author=author), True

    if ext in (".pptx", ".ppt"):
        title = Path(rel_path).stem.replace("_", " ").title()
        slides_data: Any = text_content
        if json_obj and isinstance(json_obj, dict):
            title = json_obj.get("title") or title
            slides_data = json_obj.get("slides") or json_obj.get("content") or text_content
        return generate_pptx(title, slides_data), True

    if ext in (".xlsx", ".xls"):
        sheet_name = Path(rel_path).stem.replace("_", " ").title()[:31]
        headers: list[str] = []
        rows: list[list[Any]] = []
        if json_obj and isinstance(json_obj, dict):
            sheet_name = json_obj.get("sheet_name") or sheet_name
            headers = json_obj.get("headers") or []
            rows = json_obj.get("rows") or []
        if not headers and not rows:
            headers, rows = _parse_table_data(text_content)
        return generate_xlsx(sheet_name, headers, rows), True

    if ext == ".csv":
        if json_obj and isinstance(json_obj, dict):
            headers = json_obj.get("headers") or []
            rows = json_obj.get("rows") or []
            if headers or rows:
                return generate_csv(headers, rows), False
        return text_content, False

    return text_content, False


def create_pdf(path: str | Path, title: str = "Document", content: str = "", author: str = "Synapse One") -> Path:
    """Convenience helper: build and write a PDF directly to disk."""
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_bytes(generate_pdf(title, content, author=author))
    return p


def create_docx(path: str | Path, title: str = "Document", content: str = "", author: str = "Synapse One") -> Path:
    """Convenience helper: build and write a Word (.docx) file directly to disk."""
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_bytes(generate_docx(title, content, author=author))
    return p


def create_pptx(path: str | Path, title: str = "Presentation", slides: list[dict] | str = "") -> Path:
    """Convenience helper: build and write a PPTX directly to disk."""
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_bytes(generate_pptx(title, slides))
    return p


def create_xlsx(path: str | Path, sheet_name: str = "Sheet1", rows: list[list[Any]] | None = None, headers: list[str] | None = None) -> Path:
    """Convenience helper: build and write an XLSX directly to disk."""
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    all_rows = rows or []
    if headers:
        hdr = headers
        body = all_rows
    elif all_rows:
        hdr = [str(x) for x in all_rows[0]]
        body = all_rows[1:]
    else:
        hdr = ["Column 1", "Column 2"]
        body = [["Value 1", "Value 2"]]
    p.write_bytes(generate_xlsx(sheet_name, hdr, body))
    return p


def create_csv(path: str | Path, rows: list[list[Any]] | None = None, headers: list[str] | None = None) -> Path:
    """Convenience helper: build and write a CSV directly to disk."""
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    all_rows = rows or []
    if headers:
        hdr = headers
        body = all_rows
    elif all_rows:
        hdr = [str(x) for x in all_rows[0]]
        body = all_rows[1:]
    else:
        hdr = ["Col1", "Col2"]
        body = [["Val1", "Val2"]]
    p.write_text(generate_csv(hdr, body), encoding="utf-8")
    return p
