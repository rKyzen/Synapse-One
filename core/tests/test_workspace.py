"""Phase 4 — AI Workspace tests: files, parsers, chunking, vector stores,
retrieval, vision, indexing jobs, and the facade's pipeline decisions."""

from __future__ import annotations

import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from synapse.bootstrap import Boot, create_container  # noqa: E402
from synapse.config.paths import SynapsePaths  # noqa: E402
from synapse.contracts import ModelProvider  # noqa: E402
from synapse.domain import (  # noqa: E402
    ChatRequest,
    ChatResponse,
    ModelDescriptor,
    ModelMetadata,
    ProviderKind,
)
from synapse.domain.enums import MemoryScope, ProviderState  # noqa: E402
from synapse.workspace import (  # noqa: E402
    LocalVectorStore,
    Workspace,
    WorkspaceFileManager,
)
from synapse.workspace import chunking as chunking_mod  # noqa: E402
from synapse.workspace import parsers  # noqa: E402
from synapse.workspace.config import VectorStoreSettings  # noqa: E402
from synapse.workspace.retrieval import Retrieval  # noqa: E402
from synapse.workspace.vectors import FaissVectorStore, create_vector_store  # noqa: E402
from synapse.workspace.vision import (  # noqa: E402
    VisionPipeline,
    WorkspaceVisionUnavailable,
)

try:
    import faiss  # noqa: F401
    import numpy as np  # noqa: F401

    HAVE_FAISS = True
except Exception:  # pragma: no cover
    HAVE_FAISS = False


# -- deterministic embedding provider -----------------------------------------

TERMS = ["alpha", "beta", "gamma", "database", "function", "import", "todo", "fixme", "chart"]


def fake_vector(text: str) -> list[float]:
    """Deterministic keyword-based vector in len(TERMS) dimensions."""
    lower = text.lower()
    v = [1.0 if term in lower else 0.0 for term in TERMS]
    norm = sum(x * x for x in v) ** 0.5 or 1.0
    return [x / norm for x in v]


class FakeProvider(ModelProvider):
    provider_id = "fake"
    kind = ProviderKind.LOCAL

    def __init__(self) -> None:
        self.requests: list[ChatRequest] = []

    def initialize(self) -> None:
        pass

    def list_models(self) -> list[ModelDescriptor]:
        return [ModelDescriptor(id="vision-test", provider_id="fake")]

    def to_metadata(self, descriptor: ModelDescriptor) -> ModelMetadata | None:
        return ModelMetadata(
            id=descriptor.id,
            provider_id="fake",
            kind=self.kind,
            privacy_score=1.0,
            capabilities={"vision": 1.0, "embeddings": 1.0},
        )

    def embed(self, texts: list[str], *, model: str | None = None) -> list[list[float]] | None:
        return [fake_vector(t) for t in texts]

    def chat(self, request: ChatRequest) -> ChatResponse:
        self.requests.append(request)
        return ChatResponse(
            provider_id="fake",
            model_id="vision-test",
            kind=self.kind,
            content="VISION: " + (request.messages[-1].content or "")[:60],
            raw={},
        )

    def health(self) -> bool:
        return True

    def supports(self, capability) -> bool:  # noqa: ANN001
        return True

    def shutdown(self) -> None:
        pass


@pytest.fixture
def workspace_boot(temp_paths: SynapsePaths, monkeypatch) -> tuple[Boot, FakeProvider]:
    """A Boot whose real workspace stack runs on deterministic fake vectors."""
    monkeypatch.setenv("SYNAPSE_HOME", str(temp_paths.home))
    boot = Boot(create_container(paths=temp_paths))
    fake = FakeProvider()
    boot.providers._providers["fake"] = fake
    boot.providers._states["fake"] = ProviderState.READY
    boot.registry._models["vision-test"] = ModelMetadata(
        id="vision-test",
        provider_id="fake",
        kind=ProviderKind.LOCAL,
        privacy_score=1.0,
        capabilities={"vision": 1.0, "embeddings": 1.0},
    )
    boot.workspace._settings.vision_model = "vision-test"
    return boot, fake


def make_manager(temp_paths: SynapsePaths, workspace: Workspace | None = None):
    if workspace is not None:
        return workspace.files
    from synapse.workspace.config import WorkspaceSettings

    return WorkspaceFileManager(temp_paths, WorkspaceSettings.from_config({}))


# -- file manager ---------------------------------------------------------------


def test_upload_dedupes_by_content(temp_paths):
    mgr = make_manager(temp_paths)
    a = mgr.upload("notes.txt", b"hello world")
    b = mgr.upload("notes-copy.txt", b"hello world")
    assert a.id == b.id  # same sha256 → same file
    assert a.pipeline == "document"
    blob = temp_paths.data_dir / "workspace" / "files" / f"{a.id}.txt"
    assert blob.read_bytes() == b"hello world"


def test_upload_distinct_bytes_distinct_files(temp_paths):
    mgr = make_manager(temp_paths)
    a = mgr.upload("a.txt", b"one")
    b = mgr.upload("b.txt", b"two")
    assert a.id != b.id


def test_catalog_persists_across_instances(temp_paths):
    mgr = make_manager(temp_paths)
    info = mgr.upload("x.md", b"# hi")
    mgr2 = make_manager(temp_paths)
    reloaded = mgr2.get(info.id)
    assert reloaded is not None
    assert reloaded.name == "x.md"
    assert reloaded.sha256 == info.sha256


def test_delete_removes_blob_and_catalog(temp_paths):
    mgr = make_manager(temp_paths)
    info = mgr.upload("x.py", b"print(1)")
    assert mgr.read(info.id) is not None
    assert mgr.delete(info.id) is True
    assert mgr.get(info.id) is None
    assert mgr.read(info.id) is None
    assert mgr.delete(info.id) is False


def test_pipeline_detection_by_extension(temp_paths):
    mgr = make_manager(temp_paths)
    assert mgr.upload("img.png", b"i").pipeline == "image"
    assert mgr.upload("doc.pdf", b"d").pipeline == "document"
    assert mgr.upload("code.py", b"c").pipeline == "code"
    assert mgr.upload("CODE.js", b"c").pipeline == "code"


# -- parsers ---------------------------------------------------------------------


def _minimal_pdf(text: str = "Alpha database design notes") -> bytes:
    """Handcrafted single-page PDF with a Helvetica text run."""
    content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n%s\nendstream" % (len(content), content),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref_pos = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_pos}\n%%EOF\n".encode()
    )
    return bytes(out)


def test_parse_pdf(workspace_boot):
    doc = parsers.parse(_minimal_pdf(), "pdf")
    assert len(doc.pages) == 1
    assert "Alpha database design" in doc.text


def test_parse_docx(workspace_boot):
    import docx

    real = docx.Document()
    real.add_paragraph("Alpha paragraph one")
    real.add_paragraph("Beta paragraph two")
    from io import BytesIO

    buf = BytesIO()
    real.save(buf)
    doc = parsers.parse(buf.getvalue(), "docx")
    assert "Alpha paragraph one" in doc.text
    assert "Beta paragraph two" in doc.text


def test_parse_text_fallback(workspace_boot):
    doc = parsers.parse("plain text contents".encode("utf-8"), "txt")
    assert "plain text contents" in doc.text
    # unknown extension still parses as text
    doc2 = parsers.parse("hello".encode("utf-8"), "weird")
    assert "hello" in doc2.text


# -- chunking --------------------------------------------------------------------


def test_chunk_text_respects_size():
    text = ("word " * 600)  # ~3000 chars
    chunks = chunking_mod.chunk_text(text, 1000, 100)
    assert len(chunks) >= 3
    for c in chunks:
        assert len(c) <= 1000 + 100


def test_chunk_code_preserves_line_bounds():
    lines = [f"line {i}" for i in range(100)]
    chunks = chunking_mod.chunk_code("\n".join(lines), 30, 5)
    assert len(chunks) >= 4
    for text, start, end in chunks:
        assert start >= 1
        assert end >= start
        assert "\n".join(lines[start - 1:end]) == text


# -- vector stores ----------------------------------------------------------------


def _vs_settings(provider: str) -> VectorStoreSettings:
    return VectorStoreSettings(provider=provider)


def test_local_vector_store_add_search_delete(temp_paths):
    store = LocalVectorStore(temp_paths, _vs_settings("local"))
    store.add(
        [fake_vector("alpha database"), fake_vector("beta chart")],
        [
            {"file_id": "f1", "file_name": "a.txt", "text": "alpha database doc"},
            {"file_id": "f1", "file_name": "a.txt", "text": "beta chart image"},
        ],
    )
    assert store.count() == 2
    hits = store.search(fake_vector("alpha"), k=2)
    assert hits[0].metadata["text"] == "alpha database doc"
    # filter by file
    assert store.search(fake_vector("alpha"), k=2, file_ids=["nope"]) == []
    # delete one file
    store.delete("f1")
    assert store.count() == 0
    store.add([fake_vector("alpha database")], [{"file_id": "f2", "file_name": "b", "text": "x"}])
    store.clear()
    assert store.count() == 0


def test_local_vector_store_persists(temp_paths):
    store = LocalVectorStore(temp_paths, _vs_settings("local"))
    store.add([fake_vector("alpha")], [{"file_id": "f1", "text": "a"}])
    store2 = LocalVectorStore(temp_paths, _vs_settings("local"))
    assert store2.count() == 1
    assert store2.search(fake_vector("alpha"))[0].metadata["file_id"] == "f1"


@pytest.mark.skipif(not HAVE_FAISS, reason="faiss not installed")
def test_faiss_vector_store_contributes(temp_paths):
    store = FaissVectorStore(temp_paths, _vs_settings("faiss"))
    store.add(
        [fake_vector("alpha database"), fake_vector("beta chart")],
        [{"file_id": "f1", "text": "alpha database doc"}, {"file_id": "f2", "text": "beta chart"}],
    )
    hits = store.search(fake_vector("alpha database"), k=1)
    assert hits[0].metadata["file_id"] == "f1"
    assert store.count() == 2
    store.delete("f1")
    assert store.count() == 1


@pytest.mark.skipif(not HAVE_FAISS, reason="faiss not installed")
def test_create_vector_store_prefers_configured_engine(temp_paths):
    store = create_vector_store(temp_paths, _vs_settings("faiss"))
    assert isinstance(store, FaissVectorStore)
    store2 = create_vector_store(temp_paths, _vs_settings("local"))
    assert isinstance(store2, LocalVectorStore)


@pytest.mark.skipif(not HAVE_FAISS, reason="faiss not installed")
def test_faiss_vectors_persist_across_instances(temp_paths):
    store = FaissVectorStore(temp_paths, _vs_settings("faiss"))
    store.add([fake_vector("beta chart")], [{"file_id": "f9", "text": "beta chart"}])
    reloaded = FaissVectorStore(temp_paths, _vs_settings("faiss"))
    assert reloaded.count() == 1
    hits = reloaded.search(fake_vector("beta chart"), k=1)
    assert hits and hits[0].metadata["file_id"] == "f9"


# -- retrieval -------------------------------------------------------------------


def _seeded_store(temp_paths):
    store = LocalVectorStore(temp_paths, _vs_settings("local"))
    store.add(
        [fake_vector("alpha database"), fake_vector("beta chart")],
        [{"file_id": "f1", "file_name": "a.txt", "text": "alpha database doc"},
         {"file_id": "f2", "file_name": "b.txt", "text": "beta chart summary"}],
    )
    return store


def test_retrieval_keyword_fallback_and_render(temp_paths):
    from synapse.workspace.retrieval import Retrieval

    from synapse.workspace.config import WorkspaceSettings

    class NoEmbed:
        def embed(self, texts):
            return None

    store = _seeded_store(temp_paths)
    manager = WorkspaceFileManager(temp_paths, WorkspaceSettings())
    retrieval = Retrieval(NoEmbed(), store, manager, WorkspaceSettings())
    chunks = retrieval.retrieve("alpha database", k=4)
    assert chunks and chunks[0].file_id == "f1"
    text = Retrieval.render(chunks)
    assert "[a.txt]" in text and "alpha database doc" in text


def test_scan_code_finds_markers(workspace_boot):
    boot, _ = workspace_boot
    info = boot.workspace.upload("app.py", b"def run():\n    return None\n# TODO: add retry\nprint(1)")
    matches = boot.workspace.retrieval.scan_code(["TODO", "FIXME"], [info.id])
    assert len(matches) == 1
    assert matches[0].line == 3
    assert "TODO" in matches[0].text
    assert matches[0].file_name == "app.py"


# -- vision ----------------------------------------------------------------------


def test_vision_analyze_sends_base64_image(workspace_boot):
    boot, fake = workspace_boot
    answer = boot.workspace.vision.analyze(b"PNG-BYTES", "what is this?")
    assert answer.startswith("VISION:")
    assert fake.requests
    images = fake.requests[-1].messages[-1].images
    assert images and images[0] == "UE5HLUJZVEVT"  # base64("PNG-BYTES")


def test_vision_unavailable_raises(workspace_boot):
    boot, fake = workspace_boot
    boot.providers._providers.clear()
    with pytest.raises(WorkspaceVisionUnavailable):
        boot.workspace.vision.analyze(b"x", "prompt")


# -- jobs ------------------------------------------------------------------------


def test_job_registry_progress(workspace_boot):
    boot, _ = workspace_boot
    job_id = boot.workspace.index("missing-file-id")
    job = boot.workspace.job(job_id)
    assert job is not None
    assert job["status"] in ("running", "failed")
    assert job["progress"] >= 0
    latest = boot.workspace.jobs.latest_for("missing-file-id")
    assert latest is not None


# -- facade pipeline decisions ---------------------------------------------------


def _index_document_boot(boot: Boot, text: str = "alpha database design") -> str:
    info = boot.workspace.upload("design.txt", text.encode("utf-8"))
    boot.workspace.jobs = type(boot.workspace.jobs)()  # fresh registry for clean poll
    boot.workspace.indexer.start(info.id, boot.workspace.jobs)
    import time

    for _ in range(50):
        job = boot.workspace.jobs.latest_for(info.id)
        if job and job["status"] in ("completed", "failed"):
            break
        time.sleep(0.05)
    assert job and job["status"] == "completed", job
    return info.id


def test_prepare_auto_retrieves_over_indexed_files(workspace_boot):
    boot, _ = workspace_boot
    _index_document_boot(boot, "alpha database design notes")
    context, outcome = boot.workspace.prepare("what is the alpha database?")
    assert outcome.files_used
    assert outcome.local_only is True
    assert "alpha database design notes" in context


def test_prepare_attached_image_uses_vision(workspace_boot):
    boot, fake = workspace_boot
    from io import BytesIO
    import struct

    png = b"\x89PNG\r\n\x1a\n" + b"\x00" * 64
    info = boot.workspace.upload("screenshot.png", png)
    context, outcome = boot.workspace.prepare("what does this show?", [info.id])
    assert outcome.files_attached == [info.id]
    assert outcome.vision_descriptions, "vision should run for image attachments"
    assert outcome.local_only is True
    assert outcome.files_used == [info.id]  # vision answer goes to master via prefix
    assert boot.workspace.get_file(info.id).vision_description  # memoized


def test_prepare_attached_code_scans_markers(workspace_boot):
    boot, _ = workspace_boot
    info = boot.workspace.upload("main.py", ("# TODO: refactor\n"
                                             "def main():\n"
                                             "    pass\n").encode("utf-8"))
    _index_document_boot(boot, "placeholder")  # ensure an indexed text file exists
    context, outcome = boot.workspace.prepare("find todos", [info.id])
    assert outcome.code_matches, [m.model_dump() for m in outcome.code_matches]


def test_prepare_without_files_does_not_attach(workspace_boot):
    boot, _ = workspace_boot
    context, outcome = boot.workspace.prepare("hello", None)
    assert outcome.files_attached == []
    assert outcome.files_used == []
    assert "hello" not in context  # no retrieval data


def test_workspace_delete_removes_vectors(workspace_boot):
    boot, _ = workspace_boot
    info_id = _index_document_boot(boot, "gamma beta markers")
    assert boot.workspace.vector_store.count() == 1
    boot.workspace.delete(info_id)
    assert boot.workspace.vector_store.count() == 0


# -- master agent integration ----------------------------------------------------


def test_master_integrates_workspace_files(workspace_boot):
    boot, fake = workspace_boot
    _index_document_boot(boot, "alpha database design notes")
    response = boot.master.process("summarize the alpha database notes", files=None)
    assert response.workspace is not None
    assert response.workspace.files_used, response.workspace
    assert response.workspace.local_only is True
    # attached image goes through the vision pipeline as well
    png = b"\x89PNG\r\n\x1a\n" + b"\x00" * 32
    img = boot.workspace.upload("graph.png", png)
    response2 = boot.master.process("what does this graph show?", files=[img.id])
    assert response2.workspace.vision_descriptions


# -- embedding engine (adaptive model resolution) ------------------------------


class RecordingEmbedProvider(FakeProvider):
    """FakeProvider that records the model id passed to embed()."""

    def __init__(self, installed: list[str] | None = None) -> None:
        super().__init__()
        self.installed = installed or ["all-minilm:latest"]
        self.embed_models: list[str | None] = []

    def list_models(self) -> list[ModelDescriptor]:
        return [ModelDescriptor(id=m, provider_id="fake") for m in self.installed]

    def to_metadata(self, descriptor: ModelDescriptor) -> ModelMetadata | None:
        return ModelMetadata(
            id=descriptor.id,
            provider_id="fake",
            kind=self.kind,
            privacy_score=1.0,
            capabilities={"embeddings": 1.0},
        )

    def embed(self, texts: list[str], *, model: str | None = None) -> list[list[float]] | None:
        self.embed_models.append(model)
        return [fake_vector(t) for t in texts]


class EmbedProviders:
    def __init__(self, *providers) -> None:
        self._providers = list(providers)

    def all(self):
        return self._providers


def test_embedding_engine_uses_configured_model_when_installed():
    from synapse.workspace.embeddings import EmbeddingEngine

    provider = RecordingEmbedProvider(installed=["nomic-embed-text"])
    engine = EmbeddingEngine(EmbedProviders(provider), "nomic-embed-text", batch_size=2)
    vectors = engine.embed(["alpha", "beta"])
    assert vectors is not None
    assert provider.embed_models == ["nomic-embed-text"]


def test_embedding_engine_falls_back_when_configured_model_missing():
    from synapse.workspace.embeddings import EmbeddingEngine

    provider = RecordingEmbedProvider(installed=["all-minilm:latest"])
    engine = EmbeddingEngine(EmbedProviders(provider), "nomic-embed-text", batch_size=2)
    vectors = engine.embed(["alpha", "beta"])
    assert vectors is not None
    # the configured model is not installed; None lets the provider
    # auto-resolve the installed embedding model (tier-adaptive)
    assert provider.embed_models == [None]