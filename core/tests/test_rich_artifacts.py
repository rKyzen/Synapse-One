"""Test suite for rich binary artifact generation (PDF, DOCX, PPTX, XLSX, CSV).

Verifies that genuine, openable binary files are created on disk, validated by
their respective libraries (pypdf, python-docx, python-pptx, openpyxl), and
handled seamlessly by ActionEngine, ToolRegistry, and Master Orchestrator.
"""

from __future__ import annotations

import io
import json
from pathlib import Path

import docx
import openpyxl
import pptx
import pypdf
import pytest

from synapse.actions.engine import ActionEngine
from synapse.domain.enums import RequestKind
from synapse.domain.fileops import FileAction, ValidationResult
from synapse.master.orchestrator import AIMasterOrchestrator
from synapse.master.schemas import ExecutionMode
from synapse.pipeline.tools import ToolCall, ToolRegistry
from synapse.workspace.artifacts import (
    _generate_docx_library,
    _generate_docx_pure,
    _generate_pdf_pure,
    _generate_pdf_reportlab,
    _generate_pptx_library,
    _generate_pptx_pure,
    _generate_xlsx_library,
    _generate_xlsx_pure,
    build_artifact_content,
    create_csv,
    create_docx,
    create_pdf,
    create_pptx,
    create_xlsx,
    generate_csv,
    generate_docx,
    generate_pdf,
    generate_pptx,
    generate_xlsx,
)
from synapse.workspace.manifest import parse_file_manifest
from synapse.workspace.operator import FileOperator
from synapse.workspace.review import (
    _check_csv,
    _check_docx,
    _check_pdf,
    _check_pptx,
    _check_xlsx,
    validate_file,
)


@pytest.fixture
def workspace_tmp(tmp_path: Path) -> FileOperator:
    return FileOperator(tmp_path)


# ===========================================================================
# 1. PDF Generation & Validation Tests
# ===========================================================================

def test_pdf_reportlab_generation():
    title = "Photosynthesis Overview"
    content = """# Introduction
Photosynthesis is the biological process by which plants convert light into energy.

## Key Reactions
- Light-dependent reactions (Thylakoid membrane)
- Calvin cycle (Stroma)

| Stage | Input | Output |
| Light Reactions | H2O + Light | ATP + NADPH + O2 |
| Calvin Cycle | CO2 + ATP | Glucose |
"""
    pdf_bytes = _generate_pdf_reportlab(title, content, author="Synapse AI")
    assert isinstance(pdf_bytes, bytes)
    assert pdf_bytes.startswith(b"%PDF")

    # Deep inspection with pypdf
    reader = pypdf.PdfReader(io.BytesIO(pdf_bytes))
    assert len(reader.pages) >= 1
    extracted_text = "".join(page.extract_text() for page in reader.pages)
    assert "Photosynthesis" in extracted_text or "Calvin" in extracted_text

    # Validation check
    val = _check_pdf(pdf_bytes)
    assert val.ok is True
    assert "page(s)" in " ".join(val.checks)


def test_pdf_pure_fallback_generation():
    title = "Fallback Document"
    content = """# Main Heading
- Point one
- Point two
Detailed body paragraph text goes here.
"""
    pdf_bytes = _generate_pdf_pure(title, content, author="Synapse AI")
    assert isinstance(pdf_bytes, bytes)
    assert pdf_bytes.startswith(b"%PDF")

    reader = pypdf.PdfReader(io.BytesIO(pdf_bytes))
    assert len(reader.pages) >= 1
    val = _check_pdf(pdf_bytes)
    assert val.ok is True


def test_create_pdf_convenience_helper(tmp_path: Path):
    target = tmp_path / "docs" / "summary.pdf"
    out_path = create_pdf(target, title="Project Summary", content="# Objectives\n- Complete Phase 8\n- Pass all tests")
    assert out_path.exists()
    assert out_path.stat().st_size > 500

    reader = pypdf.PdfReader(str(out_path))
    assert len(reader.pages) >= 1


# ===========================================================================
# 2. DOCX Generation & Validation Tests
# ===========================================================================

def test_docx_library_generation():
    title = "Quarterly Business Review"
    content = """# Executive Summary
Q3 performance exceeded expectations with a 24% revenue increase.

## Key Highlights
- Launched Synapse One 2.0
- Reached sub-50ms deterministic fast-path latency
- Full binary document artifact generation

| Metric | Target | Actual |
| Revenue | $1.2M | $1.45M |
| Latency | 60ms | 42ms |
| Reliability | 99.5% | 100% |
"""
    docx_bytes = _generate_docx_library(title, content, author="Synapse Finance")
    assert isinstance(docx_bytes, bytes)

    # Validate with python-docx
    doc = docx.Document(io.BytesIO(docx_bytes))
    assert len(doc.paragraphs) >= 3
    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert len(table.rows) == 4
    assert table.rows[0].cells[0].text == "Metric"

    val = _check_docx(docx_bytes)
    assert val.ok is True
    assert "docx package ok" in " ".join(val.checks)


def test_docx_pure_fallback_generation():
    title = "Pure DOCX Fallback"
    content = "# Heading 1\nParagraph content text\n# Heading 2\nAnother paragraph"
    docx_bytes = _generate_docx_pure(title, content, author="Synapse AI")
    assert isinstance(docx_bytes, bytes)

    doc = docx.Document(io.BytesIO(docx_bytes))
    assert len(doc.paragraphs) >= 2

    val = _check_docx(docx_bytes)
    assert val.ok is True


def test_create_docx_convenience_helper(tmp_path: Path):
    target = tmp_path / "reports" / "proposal.docx"
    out_path = create_docx(target, title="System Proposal", content="# Architecture\nDetailed spec.")
    assert out_path.exists()
    assert out_path.stat().st_size > 500

    doc = docx.Document(str(out_path))
    assert len(doc.paragraphs) >= 1


# ===========================================================================
# 3. PPTX Generation & Validation Tests
# ===========================================================================

def test_pptx_library_generation():
    title = "Space Exploration Roadmap"
    slides = [
        ("Mission Objectives", ["Establish permanent lunar outpost", "Deploy next-gen orbital telescope", "Robotic sample return from Europa"]),
        ("Key Milestones", ["Phase 1: Artemis Base Camp setup", "Phase 2: Deep space transport assembly", "Phase 3: Mars surface landing"]),
        ("Technology Stack", ["Reusable super-heavy launch vehicles", "In-situ resource utilization (ISRU)", "Nuclear thermal propulsion"]),
    ]
    pptx_bytes = _generate_pptx_library(title, slides)
    assert isinstance(pptx_bytes, bytes)

    prs = pptx.Presentation(io.BytesIO(pptx_bytes))
    assert len(prs.slides) == 4  # 1 title slide + 3 content slides

    val = _check_pptx(pptx_bytes)
    assert val.ok is True
    assert "4 slide(s)" in " ".join(val.checks)


def test_pptx_pure_fallback_generation():
    title = "Pure PPTX Presentation"
    slides = [
        ("Slide 1", ["Bullet A", "Bullet B"]),
        ("Slide 2", ["Bullet C", "Bullet D"]),
    ]
    pptx_bytes = _generate_pptx_pure(title, slides)
    assert isinstance(pptx_bytes, bytes)

    prs = pptx.Presentation(io.BytesIO(pptx_bytes))
    assert len(prs.slides) == 2

    val = _check_pptx(pptx_bytes)
    assert val.ok is True


def test_create_pptx_convenience_helper(tmp_path: Path):
    target = tmp_path / "slides" / "pitch.pptx"
    out_path = create_pptx(target, title="Startup Pitch", slides="# Problem\n- High latency\n# Solution\n- Synapse One")
    assert out_path.exists()
    assert out_path.stat().st_size > 500

    prs = pptx.Presentation(str(out_path))
    assert len(prs.slides) >= 2


# ===========================================================================
# 4. XLSX & CSV Generation & Validation Tests
# ===========================================================================

def test_xlsx_library_generation():
    sheet_name = "Expenses"
    headers = ["ID", "Category", "Amount", "Status"]
    rows = [
        [101, "Cloud Servers", 1250.50, "Approved"],
        [102, "Domain & SSL", 45.00, "Approved"],
        [103, "Developer Tools", 320.00, "Pending"],
    ]
    xlsx_bytes = _generate_xlsx_library(sheet_name, headers, rows)
    assert isinstance(xlsx_bytes, bytes)

    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
    assert sheet_name in wb.sheetnames
    ws = wb[sheet_name]
    assert ws.cell(row=1, column=1).value == "ID"
    assert ws.cell(row=2, column=2).value == "Cloud Servers"
    assert ws.cell(row=2, column=3).value == 1250.50

    val = _check_xlsx(xlsx_bytes)
    assert val.ok is True
    assert "1 sheet(s)" in " ".join(val.checks)


def test_xlsx_pure_fallback_generation():
    sheet_name = "Data"
    headers = ["ColA", "ColB"]
    rows = [["Val1", "100"], ["Val2", "200"]]
    xlsx_bytes = _generate_xlsx_pure(sheet_name, headers, rows)
    assert isinstance(xlsx_bytes, bytes)

    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
    assert len(wb.sheetnames) >= 1
    val = _check_xlsx(xlsx_bytes)
    assert val.ok is True


def test_csv_generation_and_validation(tmp_path: Path):
    headers = ["Name", "Role", "Score"]
    rows = [["Alice", "Architect", "98"], ["Bob", "Specialist", "95"]]
    csv_text = generate_csv(headers, rows)
    assert "Alice,Architect,98" in csv_text

    target = tmp_path / "data" / "users.csv"
    out_path = create_csv(target, headers=headers, rows=rows)
    assert out_path.exists()
    assert "Bob,Specialist,95" in out_path.read_text(encoding="utf-8")

    val = _check_csv(out_path.read_text(encoding="utf-8"))
    assert val.ok is True
    assert "3 rows" in " ".join(val.checks)


# ===========================================================================
# 5. Universal build_artifact_content Dispatcher Tests
# ===========================================================================

def test_build_artifact_content_markdown_to_pdf():
    raw_md = "# Climate Science\n- Atmospheric CO2\n- Global surface temperature"
    data, is_binary = build_artifact_content("reports/climate.pdf", raw_md)
    assert is_binary is True
    assert isinstance(data, bytes)
    assert data.startswith(b"%PDF")
    assert _check_pdf(data).ok is True


def test_build_artifact_content_markdown_to_docx():
    raw_md = "# Technical Specification\nThis document outlines the API contracts."
    data, is_binary = build_artifact_content("docs/spec.docx", raw_md)
    assert is_binary is True
    assert isinstance(data, bytes)
    assert _check_docx(data).ok is True


def test_build_artifact_content_markdown_to_pptx():
    raw_md = "# Title Slide\n- First point\n# Slide Two\n- Second point"
    data, is_binary = build_artifact_content("slides/deck.pptx", raw_md)
    assert is_binary is True
    assert isinstance(data, bytes)
    assert _check_pptx(data).ok is True


def test_build_artifact_content_csv_to_xlsx():
    raw_csv = "Product,Price,Quantity\nWidget A,19.99,100\nWidget B,29.99,50"
    data, is_binary = build_artifact_content("inventory/products.xlsx", raw_csv)
    assert is_binary is True
    assert isinstance(data, bytes)
    assert _check_xlsx(data).ok is True


def test_build_artifact_content_json_structured_to_pdf():
    raw_json = json.dumps({
        "title": "Quantum Computing",
        "author": "Dr. Researcher",
        "content": "### Qubits and Superposition\nQubits can exist in state |0>, |1>, or superposition.",
    })
    data, is_binary = build_artifact_content("research/quantum.pdf", raw_json)
    assert is_binary is True
    assert isinstance(data, bytes)
    assert _check_pdf(data).ok is True


def test_build_artifact_content_json_structured_to_xlsx():
    raw_json = json.dumps({
        "sheet_name": "Financials",
        "headers": ["Quarter", "Revenue", "Margin"],
        "rows": [["Q1", 500000, "22%"], ["Q2", 650000, "25%"]],
    })
    data, is_binary = build_artifact_content("finance/q_results.xlsx", raw_json)
    assert is_binary is True
    assert isinstance(data, bytes)
    assert _check_xlsx(data).ok is True


# ===========================================================================
# 6. ActionEngine Real File Writes & Manifest Parsing Tests
# ===========================================================================

def test_action_engine_writes_real_binary_artifacts(workspace_tmp: FileOperator):
    engine = ActionEngine(workspace_tmp)

    ops = [
        {"action": "write", "path": "docs/report.pdf", "content": "# Annual Report\n- Growth: 40%\n- Retention: 95%"},
        {"action": "write", "path": "docs/summary.docx", "content": "# Summary\nAll deliverables completed."},
        {"action": "write", "path": "slides/intro.pptx", "content": "# Synapse Intro\n- Ultra low latency\n# Features\n- Real artifacts"},
        {"action": "write", "path": "data/metrics.xlsx", "content": "Metric,Score\nSpeed,99\nQuality,100"},
        {"action": "write", "path": "data/feed.csv", "content": "id,event\n1,startup\n2,ready"},
    ]

    actions, validations = engine.apply(ops)
    assert len(actions) == 5
    assert all(a.status == "ok" for a in actions)
    assert len(validations) == 5
    assert all(v.ok for v in validations)

    # Verify files on physical disk
    pdf_path = workspace_tmp.path_for("docs/report.pdf")
    assert pdf_path.exists()
    assert pdf_path.read_bytes().startswith(b"%PDF")

    docx_path = workspace_tmp.path_for("docs/summary.docx")
    assert docx_path.exists()
    doc = docx.Document(str(docx_path))
    assert len(doc.paragraphs) >= 1

    pptx_path = workspace_tmp.path_for("slides/intro.pptx")
    assert pptx_path.exists()
    prs = pptx.Presentation(str(pptx_path))
    assert len(prs.slides) >= 2

    xlsx_path = workspace_tmp.path_for("data/metrics.xlsx")
    assert xlsx_path.exists()
    wb = openpyxl.load_workbook(str(xlsx_path))
    assert len(wb.sheetnames) >= 1

    csv_path = workspace_tmp.path_for("data/feed.csv")
    assert csv_path.exists()
    assert "startup" in csv_path.read_text(encoding="utf-8")


def test_manifest_parser_recognizes_artifacts_array():
    raw_response = """Here is the generated presentation and report:
```json
{
  "artifacts": [
    {
      "path": "pitch/deck.pptx",
      "content": "# Problem\n- Inefficiency\n# Solution\n- AI Agents"
    },
    {
      "path": "docs/brief.pdf",
      "content": "# Brief\nExecutive summary details."
    }
  ]
}
```
"""
    ops = parse_file_manifest(raw_response)
    assert len(ops) == 2
    assert ops[0]["path"] == "pitch/deck.pptx"
    assert ops[1]["path"] == "docs/brief.pdf"


def test_manifest_parser_recognizes_create_artifact_actions():
    raw_response = """
{
  "operations": [
    {
      "action": "create_pdf",
      "path": "manual/guide.pdf",
      "content": "# Guide\nStep 1: Install\nStep 2: Run"
    },
    {
      "action": "create_xlsx",
      "path": "data/log.xlsx",
      "content": "Time,Status\n10:00,OK\n11:00,OK"
    }
  ]
}
"""
    ops = parse_file_manifest(raw_response)
    assert len(ops) == 2
    assert ops[0]["action"] == "write"
    assert ops[0]["path"] == "manual/guide.pdf"
    assert ops[1]["action"] == "write"
    assert ops[1]["path"] == "data/log.xlsx"


# ===========================================================================
# 7. ToolRegistry Artifact Tool Calls Tests
# ===========================================================================

def test_tool_registry_artifact_tools(workspace_tmp: FileOperator):
    registry = ToolRegistry(file_operator=workspace_tmp)

    tools = {t.name: t for t in registry.list_tools()}
    assert "create_pdf" in tools
    assert "create_docx" in tools
    assert "create_pptx" in tools
    assert "create_xlsx" in tools
    assert "create_csv" in tools

    # Execute create_pdf
    call_pdf = ToolCall(id="t1", tool_name="create_pdf", arguments={"path": "out/doc.pdf", "title": "Tool Report", "content": "# Generated\nvia tool."})
    res_pdf = registry.execute(call_pdf)
    assert res_pdf.success is True
    assert workspace_tmp.exists("out/doc.pdf")
    assert _check_pdf(workspace_tmp.path_for("out/doc.pdf").read_bytes()).ok is True

    # Execute create_docx
    call_docx = ToolCall(id="t2", tool_name="create_docx", arguments={"path": "out/doc.docx", "title": "Word Doc", "content": "# Content\nParagraph text."})
    res_docx = registry.execute(call_docx)
    assert res_docx.success is True
    assert workspace_tmp.exists("out/doc.docx")
    assert _check_docx(workspace_tmp.path_for("out/doc.docx").read_bytes()).ok is True

    # Execute create_pptx
    call_pptx = ToolCall(id="t3", tool_name="create_pptx", arguments={"path": "out/deck.pptx", "title": "Slides", "slides_or_content": "# Slide 1\n- A\n# Slide 2\n- B"})
    res_pptx = registry.execute(call_pptx)
    assert res_pptx.success is True
    assert workspace_tmp.exists("out/deck.pptx")
    assert _check_pptx(workspace_tmp.path_for("out/deck.pptx").read_bytes()).ok is True

    # Execute create_xlsx
    call_xlsx = ToolCall(id="t4", tool_name="create_xlsx", arguments={"path": "out/sheet.xlsx", "sheet_name": "Sales", "content": "Item,Qty\nA,10\nB,20"})
    res_xlsx = registry.execute(call_xlsx)
    assert res_xlsx.success is True
    assert workspace_tmp.exists("out/sheet.xlsx")
    assert _check_xlsx(workspace_tmp.path_for("out/sheet.xlsx").read_bytes()).ok is True

    # Execute create_csv
    call_csv = ToolCall(id="t5", tool_name="create_csv", arguments={"path": "out/table.csv", "content": "Col1,Col2\nVal1,Val2"})
    res_csv = registry.execute(call_csv)
    assert res_csv.success is True
    assert workspace_tmp.exists("out/table.csv")
    assert _check_csv(workspace_tmp.read("out/table.csv")).ok is True


# ===========================================================================
# 8. Master Orchestrator Intent Analysis for Rich Artifacts
# ===========================================================================

def test_master_orchestrator_detects_pdf_request():
    analysis = AIMasterOrchestrator._fallback_analysis("Generate a 3-page PDF report about photosynthesis in docs/report.pdf")
    assert analysis.intent == "artifact_generation"
    assert analysis.domain == "document_generation"
    assert analysis.files_needed is True
    assert analysis.artifact_required is True
    assert analysis.document_processing_needed is True
    assert analysis.execution_mode == ExecutionMode.ARTIFACT_GENERATION


def test_master_orchestrator_detects_presentation_request():
    analysis = AIMasterOrchestrator._fallback_analysis("Create a PowerPoint presentation with 5 slides on space exploration in slides/space.pptx")
    assert analysis.intent == "artifact_generation"
    assert analysis.domain == "document_generation"
    assert analysis.files_needed is True
    assert analysis.artifact_required is True
    assert analysis.document_processing_needed is True


def test_master_orchestrator_detects_excel_spreadsheet_request():
    analysis = AIMasterOrchestrator._fallback_analysis("Build an Excel spreadsheet with our quarterly expenses in data/expenses.xlsx")
    assert analysis.intent == "artifact_generation"
    assert analysis.domain == "document_generation"
    assert analysis.files_needed is True
    assert analysis.artifact_required is True
    assert analysis.document_processing_needed is True


def test_master_orchestrator_detects_word_doc_request():
    analysis = AIMasterOrchestrator._fallback_analysis("Create a Word document (.docx) summarizing the meeting notes")
    assert analysis.intent == "artifact_generation"
    assert analysis.domain == "document_generation"
    assert analysis.files_needed is True
    assert analysis.artifact_required is True
    assert analysis.document_processing_needed is True


# ===========================================================================
# 9. Validation & Resilience on Corrupt Artifacts
# ===========================================================================

def test_validate_file_fails_on_corrupt_pdf():
    corrupt_pdf = b"NOT A VALID PDF CONTENT"
    res = validate_file("docs/bad.pdf", corrupt_pdf)
    assert res.ok is False
    assert "pdf header" in " ".join(res.checks) or res.error


def test_validate_file_fails_on_corrupt_pptx():
    corrupt_pptx = b"NOT A ZIP OR PPTX"
    res = validate_file("slides/bad.pptx", corrupt_pptx)
    assert res.ok is False


def test_validate_file_fails_on_corrupt_docx():
    corrupt_docx = b"NOT A ZIP OR DOCX"
    res = validate_file("docs/bad.docx", corrupt_docx)
    assert res.ok is False


def test_validate_file_fails_on_corrupt_xlsx():
    corrupt_xlsx = b"NOT A ZIP OR XLSX"
    res = validate_file("data/bad.xlsx", corrupt_xlsx)
    assert res.ok is False


# ===========================================================================
# 10. Content Completeness & Empty Artifact Rejection Tests
# ===========================================================================

def test_validate_file_fails_on_empty_pdf():
    # PDF with no real body text
    empty_pdf = generate_pdf("", "")
    res = validate_file("docs/empty.pdf", empty_pdf)
    assert res.ok is False
    assert "insufficient" in str(res.error).lower() or "pdf content length" in res.checks


def test_validate_file_fails_on_placeholder_pdf():
    placeholder_pdf = generate_pdf("Document", "...")
    res = validate_file("docs/empty.pdf", placeholder_pdf)
    assert res.ok is False


def test_validate_file_fails_on_empty_docx():
    empty_docx = generate_docx("", "")
    res = validate_file("docs/empty.docx", empty_docx)
    assert res.ok is False
    assert "insufficient" in str(res.error).lower() or "docx content length" in res.checks


def test_validate_file_fails_on_empty_pptx():
    empty_pptx = generate_pptx("", [])
    res = validate_file("slides/empty.pptx", empty_pptx)
    assert res.ok is False
    assert "insufficient" in str(res.error).lower() or "pptx slide content" in res.checks


def test_validate_file_fails_on_empty_xlsx():
    empty_xlsx = generate_xlsx("Empty", [], [])
    res = validate_file("data/empty.xlsx", empty_xlsx)
    assert res.ok is False
    assert "insufficient" in str(res.error).lower() or "xlsx cell data" in res.checks


def test_validate_file_fails_on_empty_csv():
    res = validate_file("data/empty.csv", "")
    assert res.ok is False
    assert "empty" in str(res.error).lower() or "insufficient" in str(res.error).lower()


def test_rich_pdf_content_validation_passes():
    title = "Benefits of Local AI"
    content = """# Overview of Local AI
Local AI models provide complete data privacy, zero recurring cloud API subscription costs, and ultra-low latency inference directly on user hardware.

## Key Advantages
- Complete privacy and confidential data processing
- Sub-50ms deterministic fast-path response times
- Offline reliability without internet dependencies
- Freedom from third-party vendor lock-in and rate limits
"""
    pdf_bytes = generate_pdf(title, content)
    res = validate_file("docs/local_ai.pdf", pdf_bytes)
    assert res.ok is True

    reader = pypdf.PdfReader(io.BytesIO(pdf_bytes))
    extracted = " ".join(page.extract_text() for page in reader.pages)
    assert "Benefits of Local AI" in extracted or "Overview of Local AI" in extracted
    assert "Complete privacy" in extracted or "offline reliability" in extracted.lower()


def test_rich_docx_content_validation_passes():
    title = "Benefits of Local AI"
    content = """# Executive Summary
Deploying AI locally empowers enterprises with unmatched security, reduced operational expenses, and resilient edge compute capabilities.

## Architecture Benefits
- Full project context stored in local workspace
- Zero telemetry or private code leakage
- High-throughput local model execution
"""
    docx_bytes = generate_docx(title, content)
    res = validate_file("docs/local_ai.docx", docx_bytes)
    assert res.ok is True

    doc = docx.Document(io.BytesIO(docx_bytes))
    all_text = " ".join(p.text for p in doc.paragraphs)
    assert "Deploying AI locally" in all_text


def test_rich_pptx_content_validation_passes():
    title = "Local AI Architecture"
    slides = [
        ("Why Local AI?", ["Zero cloud subscription costs", "Guaranteed offline availability", "Complete source code confidentiality"]),
        ("Performance Metrics", ["Sub-50ms response latency", "Optimized memory footprint", "Dynamic multi-tier hardware allocation"]),
    ]
    pptx_bytes = generate_pptx(title, slides)
    res = validate_file("slides/local_ai.pptx", pptx_bytes)
    assert res.ok is True


def test_rich_xlsx_content_validation_passes():
    sheet_name = "Benchmark"
    headers = ["Model", "Hardware Tier", "Latency (ms)", "Status"]
    rows = [
        ["Qwen 2.5 Coder 7B", "Tier 2 (RTX 4060)", 42.5, "Pass"],
        ["Gemma 3 4B", "Tier 1 (Integrated)", 68.0, "Pass"],
    ]
    xlsx_bytes = generate_xlsx(sheet_name, headers, rows)
    res = validate_file("data/benchmarks.xlsx", xlsx_bytes)
    assert res.ok is True


# ===========================================================================
# 11. Architect Directive: Force Substantial Task-Specific Content Tests
# ===========================================================================

def test_question_paper_meta_description_fails_validation():
    """Verify that a question paper request with only a meta-description fails validation."""
    meta_content = (
        "This document contains a comprehensive question paper on Reproduction for class 12 CBSE board, "
        "ensuring a thorough understanding of the concepts covered in the syllabus. "
        "The following sections provide questions for students."
    )
    pdf_bytes = generate_pdf("Question Paper on Reproduction", meta_content)
    res = validate_file("docs/question_paper_reproduction.pdf", pdf_bytes)
    assert res.ok is False
    assert "insufficient questions" in str(res.error).lower() or "meta-description" in str(res.error).lower()


def test_question_paper_with_real_numbered_questions_passes():
    """Verify that a question paper request with actual numbered questions passes validation."""
    real_content = """# Class 12 Biology Examination - Reproduction
**Time: 3 Hours | Total Marks: 70**

## Section A: Multiple Choice Questions
1. Which of the following organisms reproduces by multiple fission?
   (a) Amoeba  (b) Plasmodium  (c) Yeast  (d) Hydra
2. The structural and functional unit between developing embryo and maternal body is called:
   (a) Placenta  (b) Umbilical cord  (c) Amnion  (d) Chorion
3. Transfer of pollen grains from anther to stigma of another flower of same plant is:
   (a) Autogamy  (b) Geitonogamy  (c) Xenogamy  (d) Cleistogamy

## Section B: Short Answer Questions
4. Differentiate between spermatogenesis and oogenesis with three distinct physiological differences.
5. Explain the nutritive function of tapetum and protective role of endothecium in microsporangium.
6. Describe the morphological structure of human sperm with labeled parts (acrosome, middle piece, tail).
"""
    pdf_bytes = generate_pdf("Question Paper on Reproduction", real_content)
    res = validate_file("docs/question_paper_reproduction.pdf", pdf_bytes)
    assert res.ok is True
    assert "pdf valid" in res.checks


def test_short_report_meta_description_fails_validation():
    """Verify that a report containing only generic meta-description fails validation."""
    meta_content = (
        "This report provides a comprehensive overview of the Q3 performance. "
        "The following sections outline the key findings and financial highlights for the quarter."
    )
    docx_bytes = generate_docx("Quarterly Performance Report", meta_content)
    res = validate_file("docs/quarterly_report.docx", docx_bytes)
    assert res.ok is False
    assert "meta-description" in str(res.error).lower() or "report" in str(res.error).lower()


def test_short_report_with_substantive_content_passes():
    """Verify that a report with real analytical sections and data passes validation."""
    real_report = """# Q3 System Architecture & Performance Report

## 1. Executive Summary
During Q3, Synapse One achieved sub-50ms deterministic routing latency and 100% offline verification across all workspace operations. Hardware-adaptive resource scheduling reduced VRAM pressure by 34%.

## 2. Key Metrics & Benchmarks
- **Fast-Path Latency:** 42ms average response time for direct-answer intents.
- **Verification Pass Rate:** 100% across 1,050+ automated regression suites.
- **Memory Footprint:** Peak working set capped under 8.2GB on Tier 2 GPUs.

## 3. Next Steps & Q4 Roadmap
Phase 9 will introduce multi-agent concurrent DAG planning and zero-copy IPC buffers for high-bandwidth model execution.
"""
    docx_bytes = generate_docx("Quarterly Performance Report", real_report)
    res = validate_file("docs/quarterly_report.docx", docx_bytes)
    assert res.ok is True
    assert "docx package ok" in res.checks


def test_presentation_meta_description_fails_validation():
    """Verify that a PPTX with only meta-description filler fails validation."""
    meta_slides = [
        ("Overview", ["This presentation provides a comprehensive overview of AI safety.", "The following slides outline the topics."])
    ]
    pptx_bytes = generate_pptx("AI Safety Overview", meta_slides)
    res = validate_file("slides/ai_safety_presentation.pptx", pptx_bytes)
    assert res.ok is False
    assert "meta-description" in str(res.error).lower() or "insufficient" in str(res.error).lower()


def test_3_slide_presentation_with_substantive_bullets_passes():
    """Verify that a 3-slide presentation with real task-specific bullets passes validation."""
    slides = [
        ("AI Alignment Fundamentals", ["Reinforcement Learning from Human Feedback (RLHF)", "Direct Preference Optimization (DPO)", "Mechanistic Interpretability & Circuit Analysis"]),
        ("Robustness & Safety Guardrails", ["Adversarial prompt injection defense mechanisms", "Sandboxed workspace file execution with rollback", "Deterministic output verification & auto-repair"]),
        ("Deployment & Operational Security", ["Hardware-aware local model resource quotas", "Air-gapped execution with zero cloud telemetry", "Real-time action logging & audit trails"]),
    ]
    pptx_bytes = generate_pptx("AI Safety & Alignment", slides)
    res = validate_file("slides/ai_safety_presentation.pptx", pptx_bytes)
    assert res.ok is True
    assert "3 slide(s)" in " ".join(res.checks) or "4 slide(s)" in " ".join(res.checks)


def test_structured_intermediate_question_paper_dispatch():
    """Verify that specialist structured JSON questions format produces valid PDF with questions."""
    structured_data = {
        "title": "Class 12 Biology Test - Reproduction",
        "instructions": "Answer all questions. Each question carries marks as indicated.",
        "questions": [
            {"question": "What is vegetative propagation? Give two examples.", "marks": 2},
            {"question": "Explain the stages of microsporogenesis in angiosperms.", "marks": 5},
            {"question": "Differentiate between menarche and menopause.", "marks": 2},
            {"question": "Describe the structure and function of the corpus luteum.", "marks": 3},
            {"question": "What is amniocentesis? State its statutory ban rationale.", "marks": 3},
        ]
    }
    pdf_bytes, is_bin = build_artifact_content("docs/biology_test.pdf", structured_data)
    assert is_bin is True
    assert isinstance(pdf_bytes, bytes)
    assert pdf_bytes.startswith(b"%PDF")

    res = validate_file("docs/biology_test.pdf", pdf_bytes)
    assert res.ok is True


